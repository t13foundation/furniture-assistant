from __future__ import annotations

import asyncio
import io
import json
import os

import pytest
import yaml
from fastapi.testclient import TestClient
from openpyxl import Workbook
from src.furniture.rules import ConstructionRules
from src.furniture.schema import FurnitureModule, MaterialSpec, ModuleModel
from src.webapp import api, auth

from tests.conftest import sample_file


def test_upload_request_limit_bounds_streamed_body_without_content_length(monkeypatch):
    monkeypatch.setenv("FA_MAX_REQUEST_MB", "0")
    middleware = api.UploadRequestLimitMiddleware(None)
    messages = iter([{"type": "http.request", "body": b"x", "more_body": True}])
    sent = []

    async def receive():
        return next(messages)

    async def send(message):
        sent.append(message)

    async def downstream(scope, receive, send):
        await receive()

    middleware.app = downstream
    asyncio.run(
        middleware(
            {"type": "http", "method": "POST", "path": "/api/projects", "headers": []},
            receive,
            send,
        )
    )
    assert sent[0]["status"] == 413


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("FA_ENV", "development")
    monkeypatch.setenv("FA_OUTPUT_DIR", str(tmp_path / "output"))
    monkeypatch.delenv("FA_USERS_FILE", raising=False)
    with TestClient(api.app) as test_client:
        yield test_client


def _create(client: TestClient) -> str:
    response = client.post("/api/projects", files=[("files", ("cabinet.png", b"png", "image/png"))])
    assert response.status_code == 201
    return response.json()["id"]


def test_request_ids_and_json_request_log(client, monkeypatch):
    import json
    import logging
    from io import StringIO

    from src.webapp import observability

    monkeypatch.setenv("FA_LOG_FORMAT", "json")
    observability.configure_logging()
    stream = StringIO()
    handler = observability._logging_handler
    assert isinstance(handler, logging.StreamHandler)
    previous_stream = handler.setStream(stream)
    try:
        echoed = client.get("/api/health", headers={"X-Request-ID": "request-123"})
        assert echoed.headers["X-Request-ID"] == "request-123"

        generated = client.get("/api/health", headers={"X-Request-ID": "invalid id"})
        request_id = generated.headers["X-Request-ID"]
        assert request_id and request_id != "invalid id"
        assert len(request_id) == 32

        created = client.post(
            "/api/projects",
            files=[("files", ("cabinet.png", b"png", "image/png"))],
            headers={"X-Request-ID": "create-project-request"},
        )
        assert created.status_code == 201
        project_id = created.json()["id"]
        project_response = client.get(f"/api/projects/{project_id}", headers={"X-Request-ID": "project-request"})
        assert project_response.status_code == 200
    finally:
        handler.setStream(previous_stream)

    lines = [json.loads(line) for line in stream.getvalue().splitlines()]
    assert any(line["request_id"] == "request-123" for line in lines)
    assert any(line["request_id"] == request_id for line in lines)
    assert any(line["request_id"] == "create-project-request" and line["project_id"] == project_id for line in lines)
    assert any(line["request_id"] == "project-request" and line["project_id"] == project_id for line in lines)
    assert all("ts" in line and "level" in line and "logger" in line and "message" in line for line in lines)


def test_newline_upload_filename_is_logged_as_one_line(client, monkeypatch, caplog):
    import time

    from src.webapp import analysis
    from src.webapp.store import JobItem

    caplog.set_level("INFO")
    monkeypatch.setattr(analysis, "plan_items", lambda *_: [JobItem(file_index=0, page_index=0)])

    async def model_available(_request):
        return None

    monkeypatch.setattr(api, "_probe_model_server", model_available)

    def fail_extract(*_args):
        raise RuntimeError("stop after filename is logged")

    monkeypatch.setattr(analysis, "extract_marks", fail_extract)
    project_id = _create(client)
    store = api.app.state.store
    project = store.load_project(project_id)
    project.files[0].name = "cabinet\nforged.log.png"
    (store.project_dir(project_id) / "project.json").write_text(project.model_dump_json())
    assert client.post(f"/api/projects/{project_id}/analyze").status_code == 202

    deadline = time.monotonic() + 5
    while time.monotonic() < deadline:
        job = client.get(f"/api/projects/{project_id}").json()["job"]
        if job["items"] and job["items"][0]["status"] == "error":
            break
        time.sleep(0.01)
    assert job["items"][0]["status"] == "error"
    page_lines = [line for line in caplog.text.splitlines() if "Analysis page started" in line]
    assert len(page_lines) == 1
    assert r"\n" in page_lines[0]


def test_auth_disabled_and_login_protects_api(tmp_path, monkeypatch):
    monkeypatch.setenv("FA_OUTPUT_DIR", str(tmp_path / "output"))
    monkeypatch.setenv("FA_USERS_FILE", str(tmp_path / "users.yaml"))
    monkeypatch.setenv("FA_SESSION_SECRET", "test-secret")
    with TestClient(api.app) as client:
        assert client.get("/api/auth/me").json() == {"username": None, "auth_enabled": False, "role": None}
        assert _create(client)

    from src.webapp.auth import add_user

    add_user("operator", "s3cret")
    with TestClient(api.app) as client:
        assert client.get("/api/auth/me").json() == {"username": None, "auth_enabled": True, "role": None}
        assert client.get("/api/projects").status_code == 401
        assert client.get("/api/auth/me").status_code == 200
        assert client.post("/api/auth/login", json={"username": "operator", "password": "wrong"}).status_code == 401
        response = client.post("/api/auth/login", json={"username": "operator", "password": "s3cret"})
        assert response.status_code == 200
        assert response.json() == {"username": "operator", "role": "admin"}
        assert response.cookies.get("fa_session")
        assert client.get("/api/projects").status_code == 200
        assert client.post("/api/auth/logout").status_code == 200
        assert client.get("/api/projects").status_code == 401


def test_production_requires_session_secret_and_ignores_dotenv(tmp_path, monkeypatch):
    from src.webapp import auth

    monkeypatch.setenv("FA_ENV", "production")
    monkeypatch.delenv("FA_SESSION_SECRET", raising=False)
    monkeypatch.setattr(auth, "ROOT", tmp_path / "repo")
    dotenv_path = tmp_path / ".env"
    dotenv_path.write_text("FA_SESSION_SECRET=from-dotenv\n", encoding="utf-8")
    monkeypatch.setattr(api, "DOTENV_PATH", dotenv_path)

    with pytest.raises(RuntimeError, match="FA_SESSION_SECRET is required when FA_ENV=production"), TestClient(api.app):
        pass

    assert "FA_SESSION_SECRET" not in os.environ
    assert not (tmp_path / "repo" / "config" / ".session_secret").exists()


@pytest.mark.parametrize("environment", ["prod", "staging", "", "typo", " production "])
def test_invalid_fa_environment_fails_startup(monkeypatch, environment):
    monkeypatch.setenv("FA_ENV", environment)
    with pytest.raises(RuntimeError, match="FA_ENV must be 'development' or 'production'"), TestClient(api.app):
        pass


@pytest.mark.parametrize(("environment", "expected"), [("DEVELOPMENT", False), ("Production", True)])
def test_fa_environment_is_case_insensitive(tmp_path, monkeypatch, environment, expected):
    from src.webapp import auth

    monkeypatch.setenv("FA_ENV", environment)
    monkeypatch.setenv("FA_SESSION_SECRET", "test-secret")
    monkeypatch.setenv("FA_OUTPUT_DIR", str(tmp_path / "output"))
    with TestClient(api.app):
        assert auth.is_production() is expected


def test_dotenv_cannot_select_runtime_mode(tmp_path, monkeypatch):
    from src.webapp import auth

    monkeypatch.delenv("FA_ENV", raising=False)
    monkeypatch.delenv("FA_SESSION_SECRET", raising=False)
    monkeypatch.setenv("FA_OUTPUT_DIR", str(tmp_path / "output"))
    dotenv_path = tmp_path / ".env"
    dotenv_path.write_text("FA_ENV=production\nFA_SESSION_SECRET=dotenv-secret\n", encoding="utf-8")
    monkeypatch.setattr(api, "DOTENV_PATH", dotenv_path)

    with TestClient(api.app):
        assert not auth.is_production()
        assert os.environ.get("FA_ENV") is None


@pytest.mark.parametrize("value", ["tru", "", "maybe", "2", " true "])
def test_invalid_cookie_secure_fails_startup(monkeypatch, value):
    monkeypatch.setenv("FA_ENV", "development")
    monkeypatch.setenv("FA_COOKIE_SECURE", value)
    with pytest.raises(RuntimeError, match="FA_COOKIE_SECURE must be a boolean"), TestClient(api.app):
        pass


@pytest.mark.parametrize(
    ("environment", "public_url", "explicit", "expected_secure"),
    [
        ("production", None, None, True),
        ("development", "https://example.test", None, True),
        ("development", "http://example.test", None, False),
        ("production", "https://example.test", "false", False),
        ("development", "https://example.test", "true", True),
        ("development", None, "YES", True),
        ("production", None, "0", False),
    ],
)
def test_login_cookie_secure_precedence(tmp_path, monkeypatch, environment, public_url, explicit, expected_secure):
    from src.webapp.auth import add_user

    users_file = tmp_path / "users.yaml"
    monkeypatch.setenv("FA_USERS_FILE", str(users_file))
    monkeypatch.setenv("FA_ENV", environment)
    monkeypatch.setenv("FA_SESSION_SECRET", "test-secret")
    monkeypatch.setenv("FA_OUTPUT_DIR", str(tmp_path / "output"))
    if public_url is None:
        monkeypatch.delenv("FA_PUBLIC_URL", raising=False)
    else:
        monkeypatch.setenv("FA_PUBLIC_URL", public_url)
    if explicit is None:
        monkeypatch.delenv("FA_COOKIE_SECURE", raising=False)
    else:
        monkeypatch.setenv("FA_COOKIE_SECURE", explicit)
    add_user("operator", "s3cret", users_file)

    with TestClient(api.app) as client:
        response = client.post("/api/auth/login", json={"username": "operator", "password": "s3cret"})

    assert response.status_code == 200
    assert ("secure" in response.headers["set-cookie"].lower()) is expected_secure


def test_health_is_public_and_reports_model_url(tmp_path, monkeypatch):
    monkeypatch.setenv("FA_OUTPUT_DIR", str(tmp_path / "output"))
    monkeypatch.setenv("FA_USERS_FILE", str(tmp_path / "users.yaml"))
    monkeypatch.setenv("FA_SESSION_SECRET", "test-secret")
    monkeypatch.setenv("FA_API_BASE", "http://model.example/v1")
    from src.webapp.auth import add_user

    add_user("operator", "s3cret")
    with TestClient(api.app) as client:
        assert client.get("/api/projects").status_code == 401
        response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json() == {
        "status": "ok",
        "model_base_url": "http://model.example/v1",
        "projects_root_writable": True,
        "auth_enabled": True,
    }


def test_upload_limits_return_413(client, monkeypatch):
    monkeypatch.setenv("FA_MAX_UPLOAD_MB", "0")
    response = client.post("/api/projects", files=[("files", ("cabinet.png", b"png", "image/png"))])
    assert response.status_code == 413
    assert response.json()["detail"] == "File cabinet.png exceeds 0 MB"

    monkeypatch.delenv("FA_MAX_UPLOAD_MB")
    monkeypatch.setenv("FA_MAX_FILES", "1")
    two = [("files", ("a.png", b"png", "image/png")), ("files", ("b.png", b"png", "image/png"))]
    response = client.post("/api/projects", files=two)
    assert response.status_code == 413
    assert response.json()["detail"] == "At most 1 files per project"
    assert client.get("/api/projects").json() == []


def test_role_authorization_and_legacy_admin(tmp_path, monkeypatch, two_door_module):

    monkeypatch.setenv("FA_OUTPUT_DIR", str(tmp_path / "output"))
    monkeypatch.setenv("FA_USERS_FILE", str(tmp_path / "users.yaml"))
    monkeypatch.setenv("FA_SESSION_SECRET", "test-secret")
    auth.add_user("viewer", "secret", role="viewer")
    auth.add_user("editor", "secret", role="editor")
    auth.add_user("legacy", "secret")
    users_path = tmp_path / "users.yaml"
    users = yaml.safe_load(users_path.read_text())
    users["users"]["legacy"].pop("role")
    users_path.write_text(yaml.safe_dump(users))
    auth._USERS_CACHE = None
    with TestClient(api.app) as client:
        assert client.post("/api/auth/login", json={"username": "editor", "password": "secret"}).status_code == 200
        project_id = _create(client)
        module_url = f"/api/projects/{project_id}/modules/TEST"
        assert client.put(module_url, json=two_door_module.model_dump(mode="json")).status_code == 200
        client.post("/api/auth/logout")

        viewer_login = client.post("/api/auth/login", json={"username": "viewer", "password": "secret"})
        assert viewer_login.status_code == 200
        assert viewer_login.json()["role"] == "viewer"
        assert client.get("/api/auth/me").json()["role"] == "viewer"
        assert client.get("/api/projects").status_code == 200
        assert client.get(f"/api/projects/{project_id}/activity/integrity").status_code == 200
        assert client.put(module_url, json=two_door_module.model_dump(mode="json")).status_code == 403
        assert client.post(f"/api/projects/{project_id}/analyze").status_code == 403
        assert client.patch(f"/api/projects/{project_id}", json={"name": "Renamed"}).status_code == 403
        assert client.delete(f"/api/projects/{project_id}").status_code == 403
        assert (
            client.put(
                f"/api/projects/{project_id}/quote/assignments",
                json={"line_index": 0, "label": "Quote line", "module_ids": []},
            ).status_code
            == 403
        )
        client.post("/api/auth/logout")

        assert client.post("/api/auth/login", json={"username": "editor", "password": "secret"}).status_code == 200
        pricing = client.get("/api/pricing").json()
        assert client.put("/api/pricing", json=pricing).status_code == 403
        changed = two_door_module.model_dump(mode="json")
        changed["title"] = "Edited by editor"
        assert client.put(module_url, json=changed).status_code == 200
        client.post("/api/auth/logout")

        assert client.post("/api/auth/login", json={"username": "legacy", "password": "secret"}).status_code == 200
        assert client.get("/api/auth/me").json()["role"] == "admin"
        assert client.put(module_url, json=two_door_module.model_dump(mode="json")).status_code == 200


def test_analyze_rejects_unreachable_model_without_creating_job(client):
    project_id = _create(client)
    api.app.state.config.api.base_url = "http://127.0.0.1:1/v1"

    response = client.post(f"/api/projects/{project_id}/analyze")

    assert response.status_code == 503
    assert "http://127.0.0.1:1/v1" in response.json()["detail"]
    assert (api.app.state.store.project_dir(project_id) / "job.json").exists() is False


def test_analyze_queues_when_model_probe_succeeds(client, monkeypatch):
    import httpx

    project_id = _create(client)
    api.app.state.config.api.base_url = "http://model.example/v1"
    requests = []
    client_class = httpx.AsyncClient

    def handler(request):
        requests.append(request)
        return httpx.Response(200, json={"object": "list", "data": []})

    monkeypatch.setattr(
        api.httpx,
        "AsyncClient",
        lambda **kwargs: client_class(transport=httpx.MockTransport(handler), **kwargs),
    )
    response = client.post(f"/api/projects/{project_id}/analyze")

    assert response.status_code == 202
    assert response.json() == {"status": "queued"}
    assert requests[0].url == "http://model.example/v1/models"
    assert requests[0].headers["authorization"] == "Bearer lm-studio"
    assert (api.app.state.store.project_dir(project_id) / "job.json").exists()


def test_refine_rejects_unreachable_model_without_changing_module_or_audit(client, two_door_module):
    from src.furniture.schema import ModuleDraft

    project_id = _create(client)
    module_url = f"/api/projects/{project_id}/modules/TEST"
    assert client.put(module_url, json=two_door_module.model_dump(mode="json")).status_code == 200
    store = api.app.state.store
    model = store.load_module(project_id, "TEST")
    model.draft = ModuleDraft(
        module_id="TEST",
        title=two_door_module.title,
        view_bbox=two_door_module.view_bbox,
    )
    store.save_module(project_id, model)
    module_path = store.project_dir(project_id) / "modules" / "TEST.json"
    audit_path = store.project_dir(project_id) / "audit.jsonl"
    before_module = module_path.read_bytes()
    before_audit = audit_path.read_bytes()
    api.app.state.config.api.base_url = "http://127.0.0.1:1/v1"

    response = client.post(f"{module_url}/refine")

    assert response.status_code == 503
    assert "http://127.0.0.1:1/v1" in response.json()["detail"]
    assert module_path.read_bytes() == before_module
    assert audit_path.read_bytes() == before_audit


def test_total_upload_limit_rejects_before_project_creation(client, monkeypatch):
    monkeypatch.setenv("FA_MAX_REQUEST_MB", "0")
    response = client.post("/api/projects", files=[("files", ("cabinet.png", b"png", "image/png"))])
    assert response.status_code == 413
    assert response.json()["detail"] == "Request exceeds 0 MB"
    root = api.app.state.store.root
    assert not root.exists() or list(root.iterdir()) == []


def test_quote_upload_obeys_request_limit(client, monkeypatch):
    project_id = _create(client)
    quote_file = {
        "file": (
            "quote.xlsx",
            b"not a workbook",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    monkeypatch.setenv("FA_MAX_UPLOAD_MB", "0")
    oversized_file = client.post(f"/api/projects/{project_id}/quote", files=quote_file)
    assert oversized_file.status_code == 413
    assert oversized_file.json()["detail"] == "File quote.xlsx exceeds 0 MB"

    monkeypatch.delenv("FA_MAX_UPLOAD_MB")
    monkeypatch.setenv("FA_MAX_REQUEST_MB", "0")
    response = client.post(f"/api/projects/{project_id}/quote", files=quote_file)
    assert response.status_code == 413
    assert response.json()["detail"] == "Request exceeds 0 MB"


def test_audit_history_restore_and_review(client, two_door_module):
    project_id = _create(client)
    url = f"/api/projects/{project_id}/modules/TEST"
    original = two_door_module.model_dump(mode="json")
    assert client.put(url, json=original).status_code == 200
    changed = {**original, "title": "Changed"}
    assert client.put(url, json=changed).status_code == 200
    # Saving the same content again (or a cross-check pass that changes nothing) adds no version.
    assert client.post(f"{url}/answers", json={"overrides": []}).status_code == 200
    versions = client.get(f"{url}/versions").json()
    assert [entry["version"] for entry in versions] == [2, 1]
    assert client.get(f"{url}/versions/1").json()["module"]["title"] == original["title"]
    restored = client.post(f"{url}/versions/1/restore")
    assert restored.status_code == 200
    assert restored.json()["module"]["title"] == original["title"]
    activity = client.get(f"/api/projects/{project_id}/activity").json()
    integrity = client.get(f"/api/projects/{project_id}/activity/integrity").json()
    assert activity[0]["action"] == "restore"
    assert {entry["action"] for entry in activity} >= {"create_project", "editor_save", "restore"}
    assert integrity["chain_status"] == "verified"
    assert integrity["chain_error"] is None
    assert integrity["head_hash"] is not None
    exported = client.get(f"/api/projects/{project_id}/audit.jsonl")
    assert exported.status_code == 200
    assert exported.headers["content-type"].startswith("application/x-ndjson")
    review = client.get(f"/api/projects/{project_id}/review").json()
    audit_lines = exported.text.splitlines()
    audit_lines[-1] = audit_lines[-1].replace("restore", "tampered")
    (api.app.state.store.project_dir(project_id) / "audit.jsonl").write_text("\n".join(audit_lines) + "\n")
    integrity = client.get(f"/api/projects/{project_id}/activity/integrity").json()
    assert integrity["chain_status"] == "invalid"
    assert integrity["chain_error"] == f"line {len(audit_lines)}: hash mismatch"
    assert review[0]["group_id"] == two_door_module.group_id
    assert review[0]["kind"] == two_door_module.kind
    assert client.get(f"/api/projects/{project_id}/activity").status_code == 200


def test_project_pricing_markup_and_quote_comparison(client, two_door_module):
    project_id = _create(client)
    assert (
        client.put(f"/api/projects/{project_id}/modules/TEST", json=two_door_module.model_dump(mode="json")).status_code
        == 200
    )
    pricing_url = f"/api/projects/{project_id}/pricing"
    inherited = client.get(pricing_url).json()
    assert inherited["inherited"] is True
    pricing = inherited["pricing"]
    pricing["markup_percent"] = 15
    saved = client.put(pricing_url, json=pricing)
    assert saved.status_code == 200 and saved.json()["inherited"] is False
    valuation = client.get(f"/api/projects/{project_id}/valuation").json()
    assert valuation["totals"]["markup"] == round(valuation["totals"]["total_net"] * 0.15, 2)
    assert valuation["totals"]["total_with_markup"] == valuation["totals"]["total_net"] + valuation["totals"]["markup"]
    assert valuation["per_group"][0]["group_id"] == (two_door_module.group_id or "TEST")

    workbook = Workbook()
    workbook.active.append(["ZM 01.1 - LADA test", 50000])
    stream = io.BytesIO()
    workbook.save(stream)
    quote_url = f"/api/projects/{project_id}/quote"
    quote = client.post(
        quote_url,
        files={
            "file": (
                "quote.xlsx",
                stream.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert quote.status_code == 200
    payload = quote.json()
    assert payload["lines"][0]["codes"] == ["ZM011"]
    assert payload["comparison"][0]["quote_amount"] == 50000
    assert client.get(quote_url).json()["lines"] == payload["lines"]
    reverted = client.delete(pricing_url)
    assert reverted.status_code == 200 and reverted.json()["inherited"] is True


def test_user_mutations_preserve_flat_legacy_users_file(tmp_path, monkeypatch):
    users_path = tmp_path / "legacy-users.yaml"
    users_path.write_text(
        "alice:\n  salt: salt\n  password_hash: hash\n",
        encoding="utf-8",
    )
    monkeypatch.setenv("FA_USERS_FILE", str(users_path))
    auth._USERS_CACHE = None

    auth.set_user_role("alice", "viewer", users_path)
    assert auth.load_users()["alice"]["role"] == "viewer"

    auth.add_user("bob", "secret", users_path, role="editor")
    assert set(auth.load_users()) == {"alice", "bob"}
    auth.remove_user("bob", users_path)
    assert set(auth.load_users()) == {"alice"}
    saved = yaml.safe_load(users_path.read_text(encoding="utf-8"))
    assert "users" not in saved


def test_quote_manual_assignment_ignored_line_and_reimport(client, two_door_module):
    project_id = _create(client)
    assert (
        client.put(f"/api/projects/{project_id}/modules/TEST", json=two_door_module.model_dump(mode="json")).status_code
        == 200
    )
    workbook = Workbook()
    workbook.active.append(["ELEMENTY DODATKOWE", 50000])
    workbook.active.append(["POZYCJA DO POMINIĘCIA", 70000])
    stream = io.BytesIO()
    workbook.save(stream)
    upload = {
        "file": (
            "quote.xlsx",
            stream.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    url = f"/api/projects/{project_id}/quote"
    compared = client.post(url, files=upload).json()
    assert compared["comparison"][0]["match"] == "none"
    modules_url = f"/api/projects/{project_id}"
    module_id = client.get(modules_url).json()["modules"][0]["module_id"]
    assignment_url = f"{url}/assignments"
    mapped = client.put(
        assignment_url,
        json={"line_index": 0, "label": compared["lines"][0]["label"], "module_ids": [module_id]},
    )
    assert mapped.status_code == 200
    assert mapped.json()["comparison"][0]["match"] == "manual"
    assert mapped.json()["comparison"][0]["module_ids"] == [module_id]
    overlap = client.put(
        assignment_url,
        json={"line_index": 1, "label": compared["lines"][1]["label"], "module_ids": [module_id]},
    )
    assert overlap.status_code == 409

    ignored = client.put(
        assignment_url,
        json={"line_index": 1, "label": compared["lines"][1]["label"], "ignored": True},
    )
    assert ignored.status_code == 200
    assert ignored.json()["comparison"][1]["match"] == "ignored"
    assert ignored.json()["total_quote"] == 50000
    assert ignored.json()["suggested_markup_percent"] == mapped.json()["suggested_markup_percent"]

    reimported = client.post(url, files=upload)
    assert reimported.status_code == 200
    persisted = reimported.json()["comparison"]
    assert persisted[0]["match"] == "manual" and persisted[0]["module_ids"] == [module_id]
    assert persisted[1]["match"] == "ignored"


def test_empty_quote_import_is_retrievable(client):
    project_id = _create(client)
    workbook = Workbook()
    stream = io.BytesIO()
    workbook.save(stream)
    url = f"/api/projects/{project_id}/quote"
    uploaded = client.post(
        url,
        files={
            "file": (
                "empty.xlsx",
                stream.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert uploaded.status_code == 200
    assert uploaded.json()["lines"] == []
    assert client.get(url).status_code == 200
    assert client.get(url).json()["lines"] == []


def test_quote_reimport_carryover_uses_unique_label_and_amount(client):
    project_id = _create(client)
    url = f"/api/projects/{project_id}/quote"

    def workbook_upload(amounts):
        workbook = Workbook()
        for amount in amounts:
            workbook.active.append(["Transport", amount])
        stream = io.BytesIO()
        workbook.save(stream)
        return {
            "file": (
                "quote.xlsx",
                stream.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        }

    first = client.post(url, files=workbook_upload([100, 500])).json()
    ignored = client.put(
        f"{url}/assignments",
        json={"line_index": 0, "label": first["lines"][0]["label"], "ignored": True},
    )
    assert ignored.status_code == 200

    imported = client.post(url, files=workbook_upload([500, 100]))
    assert imported.status_code == 200
    rows = imported.json()["comparison"]
    assert rows[0]["quote_amount"] == 500 and rows[0]["match"] == "none"
    assert rows[1]["quote_amount"] == 100 and rows[1]["match"] == "ignored"


def test_stale_manual_module_assignment_falls_back_to_code_match():
    from types import SimpleNamespace

    from src.furniture.pricing import ModuleAmount
    from src.furniture.quote import QuoteLine, compare_quote

    model = SimpleNamespace(module=SimpleNamespace(module_id="M1", group_id="ZM01", status="ok"))
    line = QuoteLine(label="ZM01 cabinet", codes=["ZM01"], amount=120, manual_module_ids=["deleted"])
    result = compare_quote(
        [line],
        [model],
        [ModuleAmount(module_id="M1", title="Cabinet", amount=100, markup=0, total_with_markup=100)],
    )
    assert result.comparison[0].match == "code"
    assert result.comparison[0].module_ids == ["M1"]
    assert line.manual_module_ids == []


def test_users_cli_and_markup_suggestion(tmp_path, monkeypatch, capsys):
    import sys
    from types import SimpleNamespace

    from src.furniture.pricing import ModuleAmount
    from src.furniture.quote import QuoteLine, compare_quote
    from src.webapp import users

    monkeypatch.setenv("FA_USERS_FILE", str(tmp_path / "users.yaml"))
    monkeypatch.setattr(sys, "argv", ["users", "add", "operator", "--password", "secret", "--role", "viewer"])
    users.main()
    monkeypatch.setattr(sys, "argv", ["users", "set-role", "operator", "editor"])
    users.main()
    monkeypatch.setattr(sys, "argv", ["users", "list"])
    users.main()
    assert capsys.readouterr().out.strip() == "operator\teditor"
    monkeypatch.setattr(sys, "argv", ["users", "remove", "operator"])
    users.main()

    model = SimpleNamespace(module=SimpleNamespace(module_id="ZM01-1", group_id="ZM01", status="ok"))
    result = compare_quote(
        [QuoteLine(label="ZM 01.1 - Cabinet", codes=["ZM011"], amount=120)],
        [model],
        [ModuleAmount(module_id="ZM01-1", title="Cabinet", amount=100, markup=0, total_with_markup=100)],
    )
    assert result.comparison[0].app_amount == 100
    assert result.suggested_markup_percent == 20


def test_parse_real_client_quote_when_available():
    sample = sample_file("*WYCENA*.xlsx")
    if sample is None:
        pytest.skip("client quote sample is not available")
    from src.furniture.quote import parse_quote

    lines = parse_quote(sample)
    assert lines
    assert any(line.codes for line in lines)


def test_quote_line_covering_several_cabinets_is_compared_with_all_of_them(tmp_path):
    from types import SimpleNamespace

    from src.furniture.pricing import ModuleAmount
    from src.furniture.quote import compare_quote, parse_quote

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Lp.", "Opis", "Wartość"])
    # Client quotes separate code and description with an en dash.
    sheet.append([1, "ZKO2 \u2013 COFFEE POINT", 21500])
    sheet.append([2, "ZM09 ,ZM10,ZM11 \u2013 Lamele RAZEM 47,75m2", 20500])
    sheet.append([3, "ELEMENTY OZDOBNE S\u0141UP\u00d3W \u2013 4 KOMPLETY", 6000])
    workbook.save(tmp_path / "quote.xlsx")

    lines = parse_quote(tmp_path / "quote.xlsx")
    assert [line.codes for line in lines] == [["ZK02"], ["ZM09", "ZM10", "ZM11"], []]

    # The coffee point was read incompletely (priced at a fraction of its parts): it must not drive the markup.
    statuses = {"ZK02": "incomplete", "ZM09": "ok", "ZM10": "ok", "ZM11": "ok"}
    models = [SimpleNamespace(module=SimpleNamespace(module_id=m, group_id=m, status=s)) for m, s in statuses.items()]
    amounts = [ModuleAmount(module_id="ZK02", title="", amount=100, markup=0, total_with_markup=100)] + [
        ModuleAmount(module_id=m, title="", amount=1000, markup=0, total_with_markup=1000)
        for m in ("ZM09", "ZM10", "ZM11")
    ]
    result = compare_quote(lines, models, amounts)
    rows = result.comparison
    assert rows[1].app_amount == 3000 and rows[1].module_ids == ["ZM09", "ZM10", "ZM11"] and rows[1].complete
    assert not rows[0].complete
    assert rows[2].app_amount is None
    assert result.suggested_markup_percent == round((20500 / 3000 - 1) * 100, 2)
    assert result.suggestion_basis == "median quote/app ratio of 1 complete cabinet"


def test_clear_modules_archives_history_instead_of_deleting_it(client, two_door_module):
    from src.webapp.store import ProjectStore

    project_id = _create(client)
    url = f"/api/projects/{project_id}/modules/TEST"
    assert client.put(url, json=two_door_module.model_dump(mode="json")).status_code == 200
    store: ProjectStore = api.app.state.store
    with store.lock(project_id):
        store.clear_modules(project_id)
    modules_dir = store.project_dir(project_id) / "modules"
    assert not list(modules_dir.glob("*.json"))
    assert not (modules_dir / "history").exists()
    archives = [path for path in modules_dir.iterdir() if path.name.startswith("history-archive-")]
    assert len(archives) == 1 and (archives[0] / "TEST" / "1.json").is_file()
    assert client.get(f"{url}/versions").status_code == 404  # module is gone, its history archived


def test_valuation_per_cabinet_reconciles_with_the_marked_up_total(client, two_door_module):
    from src.furniture.pricing import Pricing, value_project

    def wrapped(module: FurnitureModule) -> ModuleModel:
        return ModuleModel(module=module, solids=[], cut_list=[])

    a = two_door_module.model_copy(deep=True)
    a.group_id = "ZM03"
    a.title = "WIDOK A / A VIEW"
    b = two_door_module.model_copy(deep=True)
    b.module_id = "ZM03-B"
    b.group_id = "ZM03"
    b.title = "BIUROWIEC / P6 / ZM03"
    for module in (a, b):
        module.materials = [MaterialSpec(key="M1", name="Płyta biała 18", kind="laminated_board", thickness_mm=18)]
        module.carcass_material_key = "M1"
    models = [wrapped(a), wrapped(b)]
    valuation = value_project(models, ConstructionRules(), Pricing(markup_percent=10))
    cabinet = valuation.per_group[0]
    assert cabinet.group_id == "ZM03"
    assert cabinet.title == "BIUROWIEC / P6 / ZM03"  # a view title names the view, not the cabinet
    assert cabinet.total_with_markup == pytest.approx(cabinet.amount + cabinet.markup)
    assert sum(row.total_with_markup for row in valuation.per_module) == pytest.approx(
        valuation.totals["total_with_markup"], abs=0.02
    )


def test_audit_chain_detects_tampering_removal_and_reordering(client):
    from src.webapp.audit import verify_audit

    project_id = _create(client)
    store = api.app.state.store
    for action in ("first", "second", "third"):
        store.append_audit(project_id, action, "operator")
    path = store.project_dir(project_id) / "audit.jsonl"
    original = path.read_bytes()
    assert verify_audit(path) == (True, None)

    lines = original.splitlines(keepends=True)
    changed = lines.copy()
    changed[2] = changed[2].replace(b'"second"', b'"changed"')
    path.write_bytes(b"".join(changed))
    assert verify_audit(path) == (False, "line 3: hash mismatch")

    path.write_bytes(lines[0] + lines[2] + lines[3])
    assert verify_audit(path) == (False, "line 2: prev_hash mismatch")

    path.write_bytes(b"".join(reversed(lines)))
    assert verify_audit(path) == (False, "line 1: prev_hash mismatch")


def test_audit_chain_downgrade_is_rejected_and_activity_reports_invalid(client):
    from src.webapp.audit import verify_audit

    project_id = _create(client)
    store = api.app.state.store
    store.append_audit(project_id, "first", "operator")
    store.append_audit(project_id, "second", "operator")
    path = store.project_dir(project_id) / "audit.jsonl"
    entries = [json.loads(line) for line in path.read_text().splitlines()]
    downgraded = [{key: value for key, value in entry.items() if key not in {"hash", "prev_hash"}} for entry in entries]
    path.write_text("".join(json.dumps(entry) + "\n" for entry in downgraded))

    assert verify_audit(path) == (False, "line 1: unhashed entry after chain start")
    response = client.get(f"/api/projects/{project_id}/activity")
    integrity = client.get(f"/api/projects/{project_id}/activity/integrity")
    assert response.status_code == 200 and isinstance(response.json(), list)
    assert integrity.status_code == 200
    assert integrity.json()["chain_status"] == "invalid"


def test_malformed_audit_line_is_reported_without_failing_activity(client):
    project_id = _create(client)
    store = api.app.state.store
    store.append_audit(project_id, "first", "operator")
    path = store.project_dir(project_id) / "audit.jsonl"
    path.write_bytes(path.read_bytes() + b"{not-json}\n")

    response = client.get(f"/api/projects/{project_id}/activity")
    integrity = client.get(f"/api/projects/{project_id}/activity/integrity")
    assert response.status_code == 200
    assert [entry["action"] for entry in response.json()] == ["first", "create_project"]
    assert integrity.status_code == 200
    assert integrity.json()["chain_status"] == "invalid"
    assert integrity.json()["chain_error"] == "line 3: invalid JSON"


def test_legacy_activity_is_explicitly_unverified(client):
    project_id = _create(client)
    directory = api.app.state.store.project_dir(project_id)
    path = directory / "audit.jsonl"
    path.write_text('{"at":"2026-01-01T00:00:00","user":"system","action":"legacy","module_id":null,"detail":null}\n')
    metadata_path = directory / "project.json"
    metadata = json.loads(metadata_path.read_text())
    metadata.pop("audit_chain_since")
    metadata.pop("audit_legacy_hash")
    metadata_path.write_text(json.dumps(metadata))

    integrity = client.get(f"/api/projects/{project_id}/activity/integrity")
    assert integrity.status_code == 200
    assert integrity.json()["chain_status"] == "unverified"


def test_append_terminates_valid_final_record_without_newline(client):
    from src.webapp.audit import verify_audit

    project_id = _create(client)
    store = api.app.state.store
    path = store.project_dir(project_id) / "audit.jsonl"
    original = path.read_bytes().rstrip(b"\n")
    path.write_bytes(original)

    store.append_audit(project_id, "after_unterminated_record", "operator")

    updated = path.read_bytes()
    assert updated.startswith(original + b"\n")
    assert verify_audit(path) == (True, None)
    assert json.loads(updated.splitlines()[-1])["action"] == "after_unterminated_record"


def test_partial_final_record_is_invalid_and_append_refuses_to_extend(client):
    project_id = _create(client)
    store = api.app.state.store
    path = store.project_dir(project_id) / "audit.jsonl"
    path.write_bytes(path.read_bytes() + b'{"partial":')
    original = path.read_bytes()

    with pytest.raises(ValueError, match="final line is not valid JSON"):
        store.append_audit(project_id, "must_not_append", "operator")

    assert path.read_bytes() == original
    integrity = client.get(f"/api/projects/{project_id}/activity/integrity")
    assert integrity.status_code == 200
    assert integrity.json()["chain_status"] == "invalid"
    assert integrity.json()["chain_error"] == "line 2: invalid JSON"


def test_activity_integrity_repairs_checkpoint_after_append_crash(tmp_path, monkeypatch):
    from src.webapp import store as store_module
    from src.webapp.audit import verify_audit
    from src.webapp.store import ProjectStore

    store = ProjectStore(tmp_path)
    project = store.create_project("Crash recovery", [("cabinet.png", b"png")])
    original_write_atomic = store_module._write_atomic

    def fail_checkpoint(path, text):
        if path.name == "project.json":
            raise OSError("simulated crash between audit append and checkpoint")
        original_write_atomic(path, text)

    monkeypatch.setattr(store_module, "_write_atomic", fail_checkpoint)
    with pytest.raises(OSError, match="simulated crash"):
        store.append_audit(project.id, "first", "operator")

    project_path = store.project_dir(project.id) / "project.json"
    assert json.loads(project_path.read_text())["audit_chain_since"] is None
    audit_path = store.project_dir(project.id) / "audit.jsonl"
    assert audit_path.exists()

    monkeypatch.setattr(store_module, "_write_atomic", original_write_atomic)
    assert verify_audit(audit_path) == (True, None)
    snapshot, repaired = store.audit_snapshot(project.id)
    assert repaired.audit_chain_since == 1
    assert verify_audit(snapshot, repaired.audit_chain_since, repaired.audit_legacy_hash) == (True, None)


def test_audit_chain_accepts_legacy_prefix_and_cli_verify(tmp_path, monkeypatch, capsys):
    import sys

    from src.webapp.audit import main, verify_audit
    from src.webapp.store import ProjectStore

    store = ProjectStore(tmp_path)
    project = store.create_project("Audit test", [("cabinet.png", b"png")])
    audit_path = store.project_dir(project.id) / "audit.jsonl"
    legacy = b'{"at":"2026-01-01T00:00:00","user":"system","action":"legacy","module_id":null,"detail":null}\n'
    audit_path.write_bytes(legacy)
    assert verify_audit(audit_path) == (True, None)
    store.append_audit(project.id, "after_legacy", "operator")
    assert verify_audit(audit_path) == (True, None)

    monkeypatch.setattr(sys, "argv", ["audit", "verify", str(store.project_dir(project.id))])
    assert main() == 0
    assert "valid" in capsys.readouterr().out

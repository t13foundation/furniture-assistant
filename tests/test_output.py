from pathlib import Path

from openpyxl import load_workbook
from src.output.comparison import write_review_report
from src.output.writer import write_xlsx
from src.schema.enums import ConfidenceLevel
from src.schema.extraction import EdgeBanding, ExtractionResult, FieldExtraction, Panel
from src.schema.geometry import reconcile_panel_geometry


def test_write_xlsx_consolidates_safe_panels(tmp_path: Path) -> None:
    output_path = tmp_path / "extraction_results.xlsx"
    results = [
        ExtractionResult(
            source_file="a.pdf",
            project_name=FieldExtraction(value="Project A"),
            cabinet_id=FieldExtraction(value="C-1"),
            panels=[
                Panel(
                    part_id="A-1",
                    name="Shelf",
                    length_mm=800,
                    width_mm=600,
                    thickness_mm=18,
                    board_type="MFC White",
                    material_code="K101",
                    edge_banding=EdgeBanding(top="ABS 2 mm"),
                    grain_direction="Long",
                    qty=2,
                    confidence=ConfidenceLevel.HIGH,
                ),
            ],
        ),
        ExtractionResult(
            source_file="b.pdf",
            project_name=FieldExtraction(value="Project B"),
            cabinet_id=FieldExtraction(value="C-2"),
            panels=[
                Panel(
                    name="Shelf",
                    length_mm=800,
                    width_mm=600,
                    thickness_mm=18,
                    board_type=" mfc   white ",
                    material_code=" k101 ",
                    edge_banding=EdgeBanding(top=" abs   2 MM "),
                    grain_direction=" long ",
                    qty=3,
                    confidence=ConfidenceLevel.MEDIUM,
                ),
                Panel(
                    part_id="B-2",
                    name="Edged Shelf",
                    length_mm=800,
                    width_mm=600,
                    thickness_mm=18,
                    board_type="MFC White",
                    material_code="K101",
                    edge_banding=EdgeBanding(top="ABS 1 mm"),
                    grain_direction="Long",
                    qty=1,
                    confidence=ConfidenceLevel.HIGH,
                ),
            ],
        ),
    ]

    write_xlsx(results, str(output_path))
    workbook = load_workbook(output_path, data_only=True)

    production = workbook["Production Cut List"]
    production_rows = list(production.iter_rows(min_row=2, values_only=True))
    assert len(production_rows) == 2
    consolidated = next(row for row in production_rows if row[10] == 5)
    assert consolidated[0] == "K101"
    assert consolidated[1] == "MFC White"
    assert consolidated[11] == 2.4
    assert consolidated[12] == "MEDIUM"
    assert consolidated[16] == "a.pdf#A-1; b.pdf:Shelf"

    separate_edge = next(row for row in production_rows if row[10] == 1)
    assert separate_edge[5] == "ABS 1 mm"
    assert separate_edge[11] == 0.48

    material_summary = workbook["Material Summary"]
    assert material_summary.max_row == 2
    assert material_summary["D2"].value == 6
    assert material_summary["E2"].value == 2
    assert material_summary["F2"].value == 2.88

    pricing = workbook["Pricing Overview"]
    pricing_rows = list(pricing.iter_rows(min_row=2, values_only=True))
    pricing_consolidated = next(row for row in pricing_rows if row[11] == 5)
    assert pricing_consolidated[0] == "Shelf"
    assert pricing_consolidated[12] == 2.4
    assert pricing_consolidated[13] == "MEDIUM"
    assert pricing_consolidated[14] == "READY TO PRICE"
    assert pricing_consolidated[17] == "a.pdf#A-1; b.pdf:Shelf"


def test_write_xlsx_includes_geometry_inventory(tmp_path: Path) -> None:
    output_path = tmp_path / "geometry.xlsx"
    result = ExtractionResult(
        source_file="drawing.png",
        panels=[
            Panel(name="Bok lewy", length_mm=2700, width_mm=600, qty=1),
            Panel(name="Bok prawy", length_mm=2700, width_mm=600, qty=1),
        ],
    )
    reconcile_panel_geometry(result)

    write_xlsx([result], str(output_path))
    workbook = load_workbook(output_path, data_only=True)
    rows = list(workbook["Geometry Inventory"].iter_rows(min_row=2, values_only=True))

    assert len(rows) == 2
    assert [(row[1], row[2], row[8], row[9]) for row in rows] == [
        ("G001", "side_left", 1, 1),
        ("G002", "side_right", 1, 1),
    ]


def test_write_xlsx_quarantines_unsafe_panels(tmp_path: Path) -> None:
    output_path = tmp_path / "extraction_results.xlsx"
    results = [
        ExtractionResult(
            source_file="low.pdf",
            project_name=FieldExtraction(value="Project Low"),
            cabinet_id=FieldExtraction(value="C-L"),
            panels=[
                Panel(
                    part_id="LOW-1",
                    name="Low panel",
                    length_mm=100,
                    width_mm=50,
                    thickness_mm=18,
                    board_type="MFC",
                    qty=1,
                    confidence=ConfidenceLevel.LOW,
                    notes="Check source",
                ),
            ],
        ),
        ExtractionResult(
            source_file="missing.pdf",
            project_name=FieldExtraction(value="Project Missing"),
            cabinet_id=FieldExtraction(value="C-M"),
            panels=[
                Panel(
                    part_id="MISS-1",
                    name="Missing width",
                    length_mm=100,
                    width_mm=None,
                    thickness_mm=18,
                    material_code="K101",
                    qty=1,
                    confidence=ConfidenceLevel.HIGH,
                ),
            ],
        ),
        ExtractionResult(
            source_file="qty.pdf",
            project_name=FieldExtraction(value="Project Qty"),
            cabinet_id=FieldExtraction(value="C-Q"),
            panels=[
                Panel(
                    part_id="QTY-1",
                    name="Zero quantity",
                    length_mm=100,
                    width_mm=50,
                    thickness_mm=18,
                    board_type="MFC",
                    qty=0,
                    confidence=ConfidenceLevel.HIGH,
                ),
            ],
        ),
        ExtractionResult(
            source_file="thickness.pdf",
            project_name=FieldExtraction(value="Project Thickness"),
            cabinet_id=FieldExtraction(value="C-T"),
            panels=[
                Panel(
                    part_id="THICK-1",
                    name="Front",
                    length_mm=800,
                    width_mm=600,
                    thickness_mm=None,
                    board_type="MFC",
                    qty=2,
                    confidence=ConfidenceLevel.MEDIUM,
                ),
            ],
        ),
        ExtractionResult(
            source_file="material.pdf",
            project_name=FieldExtraction(value="Project Material"),
            cabinet_id=FieldExtraction(value="C-MAT"),
            panels=[
                Panel(
                    part_id="MAT-1",
                    name="Unspecified board",
                    length_mm=800,
                    width_mm=600,
                    qty=2,
                    confidence=ConfidenceLevel.MEDIUM,
                ),
            ],
        ),
    ]

    write_xlsx(results, str(output_path))
    workbook = load_workbook(output_path, data_only=True)

    raw = workbook["Cutting List"]
    assert raw.max_row == 6
    assert [raw.cell(row=row, column=2).value for row in range(2, 7)] == [
        "LOW-1",
        "MISS-1",
        "QTY-1",
        "THICK-1",
        "MAT-1",
    ]
    assert workbook["Production Cut List"].max_row == 1
    assert workbook["Material Summary"].max_row == 1

    review = workbook["Panel Review"]
    assert review.max_row == 6
    review_rows = list(review.iter_rows(min_row=2, values_only=True))
    assert [(row[0], row[1], row[2], row[3], row[4], row[17]) for row in review_rows] == [
        ("low.pdf", "Project Low", "C-L", "LOW-1", "Low panel", "LOW confidence"),
        ("missing.pdf", "Project Missing", "C-M", "MISS-1", "Missing width", "Missing width"),
        ("qty.pdf", "Project Qty", "C-Q", "QTY-1", "Zero quantity", "Quantity must be positive"),
        ("thickness.pdf", "Project Thickness", "C-T", "THICK-1", "Front", "Missing thickness"),
        (
            "material.pdf",
            "Project Material",
            "C-MAT",
            "MAT-1",
            "Unspecified board",
            "Missing thickness; Missing board type/material code",
        ),
    ]
    pricing = workbook["Pricing Overview"]
    pricing_rows = list(pricing.iter_rows(min_row=2, values_only=True))
    assert len(pricing_rows) == 2
    missing_thickness = next(row for row in pricing_rows if row[0] == "Front")
    assert missing_thickness[3] is None
    assert missing_thickness[11] == 2
    assert missing_thickness[12] == 0.96
    assert missing_thickness[13] == "MEDIUM"
    assert missing_thickness[14] == "READY TO PRICE"
    assert missing_thickness[17] == "thickness.pdf#THICK-1"


def test_write_review_report_includes_unsafe_panels(tmp_path: Path) -> None:
    output_path = tmp_path / "review_needed.xlsx"
    result = ExtractionResult(
        source_file="unsafe.pdf",
        model_used="model",
        needs_review=True,
        review_fields=["panels"],
        panels=[
            Panel(
                name="Shelf",
                length_mm=800,
                width_mm=None,
                thickness_mm=18,
                board_type="MFC",
                confidence=ConfidenceLevel.HIGH,
            )
        ],
    )

    write_review_report([result], str(output_path))
    workbook = load_workbook(output_path, data_only=True)
    row = next(workbook.active.iter_rows(min_row=2, values_only=True))

    assert row[0:5] == ("unsafe.pdf", "model", "panel[0]", "Shelf 800.0xx18.0", "HIGH")
    assert row[5] == "Missing width"


def test_write_review_report_includes_unresolvable_source(tmp_path: Path) -> None:
    output_path = tmp_path / "review_needed.xlsx"
    result = ExtractionResult(
        source_file="cabinet.p3d",
        model_used="not-processed",
        needs_review=True,
        review_fields=["source_format"],
        unresolvable_fields=["P3D requires a Palette CAD export as PDF, DWG, or DXF before panel extraction."],
    )

    write_review_report([result], str(output_path))
    workbook = load_workbook(output_path, data_only=True)
    row = next(workbook.active.iter_rows(min_row=2, values_only=True))

    assert row == (
        "cabinet.p3d",
        "not-processed",
        "source_format",
        None,
        "N/A",
        "P3D requires a Palette CAD export as PDF, DWG, or DXF before panel extraction.",
    )

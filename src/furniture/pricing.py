from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Literal

import yaml
from openpyxl import Workbook
from pydantic import BaseModel, Field

from src.furniture.reports import material_summary, process_summary
from src.furniture.rules import ConstructionRules
from src.furniture.schema import HardwareCategory, MaterialKind, ModuleModel


class BoardPrice(BaseModel):
    match: str = ""
    kind: MaterialKind | None = None
    thickness_mm: float | None = None
    price_per_m2: float


class Pricing(BaseModel):
    currency: str = Field(default="PLN", pattern=r"^[A-Z]{3}$")  # ISO 4217 code
    boards: list[BoardPrice] = Field(default_factory=list)
    edge_banding_per_m: float | None = None
    hardware: dict[HardwareCategory, float] = Field(default_factory=dict)
    # None = labour not priced (process lines are listed as unpriced, not as free).
    labour_rate_per_hour: float | None = None
    markup_percent: float = 0
    updated_at: str | None = None


class ValuationLine(BaseModel):
    group: Literal["boards", "edges", "hardware", "processes"]
    name: str
    qty: float
    unit: str
    unit_price: float | None
    amount: float | None
    module_ids: list[str]


class ModuleAmount(BaseModel):
    module_id: str
    title: str
    amount: float  # net of markup
    markup: float
    total_with_markup: float


class GroupAmount(BaseModel):
    group_id: str
    title: str
    amount: float  # net of markup
    markup: float
    total_with_markup: float
    module_ids: list[str]


class Valuation(BaseModel):
    currency: str
    lines: list[ValuationLine]
    totals: dict[str, float]
    unpriced: int
    per_module: list[ModuleAmount]
    per_group: list[GroupAmount]
    notes: list[str]


def load_pricing(path: str | Path = "config/pricing.yaml") -> Pricing:
    file = Path(path)
    if not file.exists():
        return Pricing()
    return Pricing.model_validate(yaml.safe_load(file.read_text()) or {})


def save_pricing(pricing: Pricing, path: str | Path = "config/pricing.yaml") -> Pricing:
    updated = pricing.model_copy(update={"updated_at": datetime.now().isoformat(timespec="seconds")})
    file = Path(path)
    file.parent.mkdir(parents=True, exist_ok=True)
    file.write_text(
        yaml.safe_dump(updated.model_dump(mode="json"), allow_unicode=True, sort_keys=False), encoding="utf-8"
    )
    return updated


def _weights(per_module: dict[str, float]) -> dict[str, float]:
    """Shares of a line's amount by each module's quantity."""
    total = sum(per_module.values())
    return {module_id: qty / total for module_id, qty in per_module.items()} if total else {}


def _group_title(group_id: str, group: list[ModuleModel]) -> str:
    """A cabinet's display name: a view title such as "WIDOK A / A VIEW" names the view, not the cabinet."""
    compact = re.sub(r"[^A-Z0-9]", "", group_id.upper())
    titled = [m.module.title for m in group]
    for title in titled:
        if compact and compact in re.sub(r"[^A-Z0-9]", "", title.upper()):
            return title
    for title in titled:
        if not re.match(r"^widok\b", title.strip(), re.I):
            return title
    return group_id


def value_project(models: list[ModuleModel], rules: ConstructionRules, pricing: Pricing) -> Valuation:
    materials = material_summary(models, rules)
    processes = process_summary(models)
    lines: list[ValuationLine] = []
    module_amounts = {model.module.module_id: 0.0 for model in models}
    totals = {name: 0.0 for name in ("boards", "edges", "hardware", "processes", "total_net")}
    unpriced = 0

    def append(
        group: Literal["boards", "edges", "hardware", "processes"],
        name: str,
        qty: float,
        unit: str,
        price: float | None,
        module_ids: list[str],
        shares: dict[str, float] | None = None,
    ) -> None:
        nonlocal unpriced
        amount = round(qty * price, 2) if price is not None else None
        lines.append(
            ValuationLine(
                group=group,
                name=name,
                qty=round(qty, 4),
                unit=unit,
                unit_price=price,
                amount=amount,
                module_ids=module_ids,
            )
        )
        if amount is None:
            unpriced += 1
            return
        totals[group] += amount
        if shares:
            for module_id, share in shares.items():
                module_amounts[module_id] += amount * share
        elif module_ids:
            share = amount / len(module_ids)
            for module_id in module_ids:
                module_amounts[module_id] += share

    for row in materials.boards:
        module_areas: dict[str, float] = {}
        kind = None
        for model in models:
            if model.module.module_id not in row.module_ids:
                continue
            material = next((m for m in model.module.materials if m.name == row.material), None)
            kind = kind or (material.kind if material else None)
            area = sum(
                (panel.length_mm or 0) * (panel.width_mm or 0) * panel.qty / 1_000_000
                for panel in model.cut_list
                if panel.board_type == row.material and (panel.thickness_mm or 0) == row.thickness_mm
            )
            if area:
                module_areas[model.module.module_id] = area
        price = next(
            (
                candidate
                for candidate in pricing.boards
                if candidate.match
                and row.material
                and candidate.match.casefold() in row.material.casefold()
                and candidate.kind in (None, kind)
                and candidate.thickness_mm in (None, row.thickness_mm)
            ),
            None,
        )
        if price is None:
            price = next(
                (
                    candidate
                    for candidate in pricing.boards
                    if candidate.kind == kind
                    and candidate.kind is not None
                    and (candidate.thickness_mm is None or candidate.thickness_mm == row.thickness_mm)
                ),
                None,
            )
        if price is None:
            price = next(
                (
                    candidate
                    for candidate in pricing.boards
                    if candidate.thickness_mm == row.thickness_mm
                    and candidate.kind in (None, kind)
                    and candidate.thickness_mm is not None
                ),
                None,
            )
        unit_price = price.price_per_m2 if price else None
        area_per_sheet = rules.sheet_length_mm * rules.sheet_width_mm / 1_000_000
        bought_area = row.sheets * area_per_sheet
        shares = (
            {key: value / sum(module_areas.values()) for key, value in module_areas.items()}
            if module_areas and sum(module_areas.values())
            else {}
        )
        append(
            "boards",
            f"{row.material or 'Unspecified'} {row.thickness_mm:g} mm",
            bought_area,
            "m²",
            unit_price,
            row.module_ids,
            shares,
        )
    for row in materials.edges:
        append(
            "edges",
            row.material,
            row.length_m,
            "mb",
            pricing.edge_banding_per_m,
            row.module_ids,
            _weights(row.per_module),
        )
    for row in materials.hardware:
        price = pricing.hardware.get(row.category)
        append(
            "hardware",
            f"{row.name} {row.spec}".strip(),
            row.qty,
            row.unit,
            price,
            row.module_ids,
            _weights(row.per_module),
        )
    unplaced = [part for row in materials.boards for part in row.unplaced]
    for line in processes.lines:
        append("processes", line.operation, line.minutes / 60, "h", pricing.labour_rate_per_hour, [line.module_id])
    totals = {key: round(value, 2) for key, value in totals.items()}
    totals["total_net"] = round(sum(totals[key] for key in ("boards", "edges", "hardware", "processes")), 2)
    totals["markup"] = round(totals["total_net"] * pricing.markup_percent / 100, 2)
    totals["total_with_markup"] = round(totals["total_net"] + totals["markup"], 2)
    share = pricing.markup_percent / 100

    def with_markup(amount: float) -> tuple[float, float]:
        markup = round(amount * share, 2)
        return markup, round(amount + markup, 2)

    per_module = []
    for model in models:
        amount = round(module_amounts[model.module.module_id], 2)
        markup, total = with_markup(amount)
        per_module.append(
            ModuleAmount(
                module_id=model.module.module_id,
                title=model.module.title,
                amount=amount,
                markup=markup,
                total_with_markup=total,
            )
        )
    groups: dict[str, list[ModuleModel]] = {}
    for model in models:
        group_id = model.module.group_id or model.module.module_id
        groups.setdefault(group_id, []).append(model)
    amount_by_id = {row.module_id: row for row in per_module}
    per_group = []
    for group_id, group in groups.items():
        amount = round(sum(amount_by_id[m.module.module_id].amount for m in group), 2)
        markup = round(sum(amount_by_id[m.module.module_id].markup for m in group), 2)
        per_group.append(
            GroupAmount(
                group_id=group_id,
                title=_group_title(group_id, group),
                amount=amount,
                markup=markup,
                total_with_markup=round(amount + markup, 2),
                module_ids=[m.module.module_id for m in group],
            )
        )
    return Valuation(
        currency=pricing.currency,
        lines=lines,
        totals=totals,
        unpriced=unpriced,
        per_module=per_module,
        per_group=per_group,
        notes=[
            "Cennik zawiera przykładowe stawki i wymaga potwierdzenia.",
            *([f"Nie wyceniono {unpriced} pozycji; nie są ujęte w sumach."] if unpriced else []),
            *(
                [f"Elementy dłuższe niż arkusz (nie do rozkroju, nie wycenione): {', '.join(unplaced)}."]
                if unplaced
                else []
            ),
        ],
    )


def write_valuation_xlsx(valuation: Valuation, models: list[ModuleModel], path: str | Path) -> Path:
    workbook = Workbook()
    valuation_sheet = workbook.active
    valuation_sheet.title = "Valuation"
    valuation_sheet.append(["Group", "Name", "Quantity", "Unit", "Unit price", "Amount", "Modules"])
    for line in valuation.lines:
        valuation_sheet.append(
            [line.group, line.name, line.qty, line.unit, line.unit_price, line.amount, ", ".join(line.module_ids)]
        )
    valuation_sheet.append([])
    valuation_sheet.append(["Currency", valuation.currency])
    for key, amount in valuation.totals.items():
        valuation_sheet.append([key, amount])
    material_sheet = workbook.create_sheet("Materials")
    material_sheet.append(["Module", "Part", "Quantity", "Length mm", "Width mm", "Thickness mm", "Material"])
    for model in models:
        for panel in model.cut_list:
            material_sheet.append(
                [
                    model.module.module_id,
                    panel.name,
                    panel.qty,
                    panel.length_mm,
                    panel.width_mm,
                    panel.thickness_mm,
                    panel.board_type,
                ]
            )
    hardware_sheet = workbook.create_sheet("Hardware")
    hardware_sheet.append(["Module", "Category", "Name", "Spec", "Quantity", "Unit"])
    for model in models:
        for item in model.hardware:
            hardware_sheet.append([model.module.module_id, item.category, item.name, item.spec, item.qty, item.unit])
    process_sheet = workbook.create_sheet("Processes")
    process_sheet.append(["Module", "Process", "Operation", "Quantity", "Unit", "Minutes"])
    for model in models:
        for line in model.processes:
            process_sheet.append(
                [model.module.module_id, line.process, line.operation, line.qty, line.unit, line.minutes]
            )
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output

from __future__ import annotations

import re
from collections.abc import Collection
from pathlib import Path
from statistics import median
from typing import Any, BinaryIO

from openpyxl import load_workbook
from pydantic import BaseModel


class QuoteLine(BaseModel):
    label: str
    # Cabinet codes the line prices (one line may cover several, e.g. "ZM09, ZM10, ZM11, ZM12 - Lamele").
    codes: list[str]
    amount: float


class QuoteMatch(BaseModel):
    codes: list[str]
    label: str
    quote_amount: float
    app_amount: float | None
    ratio: float | None
    module_ids: list[str]
    # Every matched module is fully dimensioned and priced (otherwise the app prices only part of the cabinet).
    complete: bool


class QuoteComparison(BaseModel):
    lines: list[QuoteLine]
    comparison: list[QuoteMatch]
    total_quote: float
    total_app: float
    suggested_markup_percent: float
    # How the suggestion was derived, e.g. "median quote/app ratio of 3 complete cabinets".
    suggestion_basis: str


# Cabinet codes: "ZM05", "ZM 01.1", "ZK-04" (not "RAZEM 47").
_CODE = re.compile(r"\b[A-Z]{2,4} ?-? ?\d{2,3}(?:\.\d{1,2})?\b", re.I)
# Two-letter prefix with the letter O typed for a zero: "ZKO2" → "ZK02".
_ZERO_TYPO = re.compile(r"\b([A-Z]{2})O(\d)", re.I)


def normalize_code(value: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", value.upper())


def _codes(text: str) -> list[str]:
    return list(
        dict.fromkeys(normalize_code(match.group()) for match in _CODE.finditer(_ZERO_TYPO.sub(r"\g<1>0\2", text)))
    )


def parse_quote(source: str | Path | BinaryIO) -> list[QuoteLine]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    lines: list[QuoteLine] = []
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                if len(row) < 2:
                    continue
                texts = [str(value).strip() for value in row if isinstance(value, str) and value.strip()]
                if not texts:
                    continue
                label = next((text for text in texts if _codes(text)), "")
                codes = _codes(label) if label else []
                if not label:
                    label = next(
                        (
                            text
                            for text in texts
                            if len(text) > 3 and not re.match(r"^(?:lp\.|opis|wartość|razem\b)", text, re.I)
                        ),
                        "",
                    )
                amount = next((value for value in reversed(row) if isinstance(value, (int, float))), None)
                if not label or amount is None or label.casefold().startswith("razem"):
                    continue
                lines.append(QuoteLine(label=label, codes=codes, amount=float(amount)))
    finally:
        workbook.close()
    return lines


def compare_quote(
    lines: list[QuoteLine], models: list[Any], per_module: list[Any], unpriced_module_ids: Collection[str] = ()
) -> QuoteComparison:
    """Match quote lines to cabinets; ``unpriced_module_ids`` have valuation lines without a price."""
    amounts = {row.module_id: row.amount for row in per_module}
    codes: dict[str, list[Any]] = {}
    for model in models:
        module = model.module
        code = normalize_code(module.group_id or module.module_id)
        codes.setdefault(code, []).append(model)
        # A quote code may name an individual run, while the app group identifies its cabinet.
        codes.setdefault(normalize_code(module.module_id), []).append(model)
    comparison: list[QuoteMatch] = []
    total_app = 0.0
    for line in lines:
        matched = [model for code in line.codes for model in codes.get(code, [])]
        # Avoid modules inserted twice when module_id and group_id normalize identically.
        matched = list({model.module.module_id: model for model in matched}.values())
        app_amount = round(sum(amounts.get(model.module.module_id, 0.0) for model in matched), 2) if matched else None
        ratio = line.amount / app_amount if app_amount else None
        if app_amount is not None:
            total_app += app_amount
        comparison.append(
            QuoteMatch(
                codes=line.codes,
                label=line.label,
                quote_amount=line.amount,
                app_amount=app_amount,
                ratio=round(ratio, 4) if ratio is not None else None,
                module_ids=[model.module.module_id for model in matched],
                complete=bool(matched)
                and all(
                    model.module.status == "ok" and model.module.module_id not in unpriced_module_ids
                    for model in matched
                ),
            )
        )
    # Median ratio: one badly read cabinet (a desk priced at a tenth of its parts) must not set the markup.
    priced = [row for row in comparison if row.ratio is not None]
    complete = [row for row in priced if row.complete]
    basis_rows, basis = (complete, "complete") if complete else (priced, "matched, incomplete")
    ratios = [row.ratio for row in basis_rows if row.ratio is not None]
    suggested = (median(ratios) - 1) * 100 if ratios else 0.0
    suggestion_basis = (
        f"median quote/app ratio of {len(ratios)} {basis} cabinet{'s' if len(ratios) != 1 else ''}"
        if ratios
        else "no quote line matches a priced cabinet"
    )
    return QuoteComparison(
        lines=lines,
        comparison=comparison,
        total_quote=round(sum(line.amount for line in lines), 2),
        total_app=round(total_app, 2),
        suggested_markup_percent=round(suggested, 2),
        suggestion_basis=suggestion_basis,
    )

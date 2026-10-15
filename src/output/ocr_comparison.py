"""XLSX report for OCR backend comparison."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from src.output.comparison import COMPARISON_FIELDS
from src.pipeline.comparator import ComparisonResult

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
AGREE_FILL = PatternFill(start_color="6BCB77", end_color="6BCB77", fill_type="solid")
PARTIAL_FILL = PatternFill(start_color="FFD93D", end_color="FFD93D", fill_type="solid")
DISAGREE_FILL = PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid")
MISSING_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
WRAP_ALIGN = Alignment(wrap_text=True, vertical="top")


def write_ocr_comparison_report(
    comparisons: list[ComparisonResult],
    output_path: str,
) -> Path:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    _write_field_comparison(wb, comparisons)
    _write_agreement_matrix(wb, comparisons)
    _write_drawing_summary(wb, comparisons)
    _write_resource_comparison(wb, comparisons)

    wb.save(str(path))
    return path


def _write_field_comparison(wb: Workbook, comparisons: list[ComparisonResult]) -> None:
    ws = wb.active
    ws.title = "Field Comparison"

    for comp in comparisons:
        backend_names = [br.backend_name for br in comp.backends]

        ws.append([f"=== {comp.drawing_name} ==="])
        ws.append([])

        headers = ["Field"]
        for name in backend_names:
            headers.extend([f"{name} (Value)", f"{name} (Confidence)"])
        ws.append(headers)
        for cell in ws[ws.max_row]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT

        for field_name in COMPARISON_FIELDS:
            row = [field_name]
            for br in comp.backends:
                f = getattr(br.result, field_name, None)
                if f and f.value:
                    val = f.value
                    if f.value_eng and f.value_eng != f.value:
                        val += f"\n(en: {f.value_eng})"
                    row.append(val)
                    row.append(f.confidence.value)
                else:
                    row.append("")
                    row.append("")
            ws.append(row)

        ws.append([])


def _write_agreement_matrix(wb: Workbook, comparisons: list[ComparisonResult]) -> None:
    ws = wb.create_sheet("Agreement Matrix")

    for comp in comparisons:
        backend_names = [br.backend_name for br in comp.backends]

        ws.append([f"=== {comp.drawing_name} ==="])
        ws.append([])

        headers = ["Field", "Agreement", "Agree", "Disagree", "Missing"]
        for name in backend_names:
            headers.append(f"{name} Value")
        ws.append(headers)
        for cell in ws[ws.max_row]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT

        for agr in comp.field_agreements:
            ratio = agr.agreement_ratio
            if agr.missing_count == agr.total_backends:
                status = "NO DATA"
                fill = MISSING_FILL
            elif ratio >= 0.8:
                status = "FULL"
                fill = AGREE_FILL
            elif ratio >= 0.5:
                status = "PARTIAL"
                fill = PARTIAL_FILL
            else:
                status = "DISAGREE"
                fill = DISAGREE_FILL

            row = [
                agr.field_name,
                status,
                agr.agree_count,
                agr.disagree_count,
                agr.missing_count,
            ]
            for name in backend_names:
                row.append(agr.values.get(name, ""))

            ws.append(row)
            ws.cell(row=ws.max_row, column=2).fill = fill

        ws.append([])


def _write_drawing_summary(wb: Workbook, comparisons: list[ComparisonResult]) -> None:
    ws = wb.create_sheet("Per-Drawing Summary")

    headers = [
        "Drawing",
        "Type",
        "Backend",
        "Fields Extracted",
        "Needs Review",
        "Review Fields",
        "Tokens",
        "Time (s)",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for comp in comparisons:
        for br in comp.backends:
            fields_found = sum(1 for fn in COMPARISON_FIELDS if (f := getattr(br.result, fn, None)) and f.value)
            ws.append(
                [
                    comp.drawing_name,
                    comp.drawing_type.value,
                    br.backend_name,
                    f"{fields_found}/{len(COMPARISON_FIELDS)}",
                    "YES" if br.result.needs_review else "NO",
                    ", ".join(br.result.review_fields),
                    br.result.tokens_used,
                    f"{comp.timing.get(br.backend_name, 0):.1f}",
                ]
            )


def _write_resource_comparison(wb: Workbook, comparisons: list[ComparisonResult]) -> None:
    ws = wb.create_sheet("Resources")

    all_backends: set[str] = set()
    for comp in comparisons:
        for br in comp.backends:
            all_backends.add(br.backend_name)

    headers = ["Backend", "Drawings", "Total Tokens", "Total Time (s)", "Avg Time (s)", "Avg Fields"]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for name in sorted(all_backends):
        total_tokens = 0
        total_time = 0.0
        drawing_count = 0
        total_fields = 0

        for comp in comparisons:
            for br in comp.backends:
                if br.backend_name == name:
                    total_tokens += br.result.tokens_used
                    total_time += comp.timing.get(name, 0.0)
                    drawing_count += 1
                    total_fields += sum(
                        1 for fn in COMPARISON_FIELDS if (f := getattr(br.result, fn, None)) and f.value
                    )

        avg_time = total_time / drawing_count if drawing_count else 0
        avg_fields = total_fields / drawing_count if drawing_count else 0

        ws.append(
            [
                name,
                drawing_count,
                total_tokens,
                f"{total_time:.1f}",
                f"{avg_time:.1f}",
                f"{avg_fields:.1f}",
            ]
        )

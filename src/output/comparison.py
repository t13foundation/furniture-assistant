from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from src.schema.extraction import ExtractionResult
from src.schema.validation import panel_review_reasons

COMPARISON_FIELDS = [
    "project_name",
    "client",
    "cabinet_id",
    "scale",
    "units",
]

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
LOW_FILL = PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid")
MED_FILL = PatternFill(start_color="FFD93D", end_color="FFD93D", fill_type="solid")
HIGH_FILL = PatternFill(start_color="6BCB77", end_color="6BCB77", fill_type="solid")


def write_benchmark_report(
    results_by_model: dict[str, list[ExtractionResult]],
    output_path: str,
    timing_by_model: dict[str, float] | None = None,
    token_usage_by_model: dict[str, int] | None = None,
) -> Path:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    _write_comparison_sheet(wb, results_by_model)
    _write_confidence_sheet(wb, results_by_model)
    _write_resource_sheet(wb, results_by_model, timing_by_model, token_usage_by_model)

    wb.save(str(path))
    return path


def write_review_report(
    all_results: list[ExtractionResult],
    output_path: str,
) -> Path:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Review Needed"

    ws.append(
        [
            "Source File",
            "Model",
            "Field",
            "Value",
            "Confidence",
            "Notes",
        ]
    )

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for result in all_results:
        for field_name in COMPARISON_FIELDS:
            f = getattr(result, field_name, None)
            if f and f.confidence.value == "LOW":
                ws.append(
                    [
                        Path(result.source_file).name,
                        result.model_used,
                        field_name,
                        f.value or "",
                        f.confidence.value,
                        f.notes,
                    ]
                )
        if result.needs_review:
            for field_name in result.review_fields:
                if field_name == "panels":
                    for index, panel in enumerate(result.panels):
                        reasons = panel_review_reasons(panel)
                        if not reasons:
                            continue
                        dimensions = f"{panel.length_mm or ''}x{panel.width_mm or ''}x{panel.thickness_mm or ''}"
                        ws.append(
                            [
                                Path(result.source_file).name,
                                result.model_used,
                                f"panel[{index}]",
                                f"{panel.name or 'unnamed'} {dimensions}",
                                panel.confidence.value,
                                "; ".join(reasons),
                            ]
                        )
                    continue
                if field_name not in COMPARISON_FIELDS:
                    ws.append(
                        [
                            Path(result.source_file).name,
                            result.model_used,
                            field_name,
                            "",
                            "N/A",
                            "; ".join(result.unresolvable_fields) or result.notes,
                        ]
                    )
                    continue
                f = getattr(result, field_name, None)
                if f and f.confidence.value != "LOW":
                    ws.append(
                        [
                            Path(result.source_file).name,
                            result.model_used,
                            field_name,
                            f.value or "",
                            f.confidence.value,
                            f.notes,
                        ]
                    )

    wb.save(str(path))
    return path


def write_consolidated_summary(
    results_by_model: dict[str, list[ExtractionResult]],
    output_path: str,
) -> Path:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    drawings: dict[str, dict[str, ExtractionResult]] = {}
    for model_name, results in results_by_model.items():
        for r in results:
            key = Path(r.source_file).name
            drawings.setdefault(key, {})[model_name] = r

    wb = Workbook()
    ws = wb.active
    ws.title = "Best per Drawing"

    ws.append(["Drawing", "Best Model", "Fields Extracted", "Needs Review"])
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for drawing_name, model_results in sorted(drawings.items()):
        best_model = max(
            model_results.items(),
            key=lambda x: sum(1 for fn in COMPARISON_FIELDS if (f := getattr(x[1], fn, None)) and f.value),
        )
        model_name, result = best_model
        fields_filled = sum(1 for fn in COMPARISON_FIELDS if (f := getattr(result, fn, None)) and f.value)
        ws.append(
            [
                drawing_name,
                model_name,
                f"{fields_filled}/{len(COMPARISON_FIELDS)}",
                "YES" if result.needs_review else "NO",
            ]
        )

    wb.save(str(path))
    return path


def _write_comparison_sheet(
    wb: Workbook,
    results_by_model: dict[str, list[ExtractionResult]],
) -> None:
    ws = wb.active
    ws.title = "Field Comparison"

    models = list(results_by_model.keys())
    headers = ["Drawing"]
    for model in models:
        headers.extend([f"{model} (Value)", f"{model} (Confidence)"])
    ws.append(headers)

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    drawings = _align_by_drawing(results_by_model)

    for drawing_name, model_results in sorted(drawings.items()):
        row = [drawing_name]
        for model in models:
            result = model_results.get(model)
            if result:
                values = []
                confidences = []
                for fn in COMPARISON_FIELDS:
                    f = getattr(result, fn, None)
                    values.append(f.value or "")
                    confidences.append(f.confidence.value if f else "")
                row.append(" | ".join(values))
                row.append(" | ".join(confidences))
            else:
                row.extend(["", ""])
        ws.append(row)


def _write_confidence_sheet(
    wb: Workbook,
    results_by_model: dict[str, list[ExtractionResult]],
) -> None:
    ws = wb.create_sheet("Confidence Heatmap")

    models = list(results_by_model.keys())
    headers = ["Field", *models]
    ws.append(headers)

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for field_name in COMPARISON_FIELDS:
        row = [field_name]
        for model in models:
            results = results_by_model[model]
            confidences = [
                getattr(r, field_name, None).confidence.value for r in results if getattr(r, field_name, None)
            ]
            if confidences:
                high = confidences.count("HIGH") / len(confidences)
                row.append(
                    f"H:{high:.0%} M:{confidences.count('MEDIUM') / len(confidences):.0%} L:{confidences.count('LOW') / len(confidences):.0%}"
                )
            else:
                row.append("N/A")
        ws.append(row)


def _write_resource_sheet(
    wb: Workbook,
    results_by_model: dict[str, list[ExtractionResult]],
    timing_by_model: dict[str, float] | None,
    token_usage_by_model: dict[str, int] | None,
) -> None:
    ws = wb.create_sheet("Resources")
    ws.append(["Model", "Drawings Processed", "Total Tokens", "Wall Time (s)", "Avg Time/Drawing (s)"])

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for model, results in results_by_model.items():
        tokens = (
            token_usage_by_model.get(model, sum(r.tokens_used for r in results))
            if token_usage_by_model
            else sum(r.tokens_used for r in results)
        )
        wall_time = timing_by_model.get(model, 0.0) if timing_by_model else 0.0
        avg_time = wall_time / len(results) if results else 0.0

        ws.append(
            [
                model,
                len(results),
                tokens,
                f"{wall_time:.1f}",
                f"{avg_time:.1f}",
            ]
        )


def _align_by_drawing(
    results_by_model: dict[str, list[ExtractionResult]],
) -> dict[str, dict[str, ExtractionResult]]:
    drawings: dict[str, dict[str, ExtractionResult]] = {}
    for model_name, results in results_by_model.items():
        for r in results:
            key = Path(r.source_file).name
            drawings.setdefault(key, {})[model_name] = r
    return drawings

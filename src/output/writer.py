from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from src.pipeline.tokens import DrawingMetrics, SchemaOverhead
from src.schema.enums import ConfidenceLevel
from src.schema.extraction import ExtractionResult, Panel
from src.schema.validation import panel_review_reasons


def write_json(result: ExtractionResult, output_dir: str) -> Path:
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_model = result.model_used.replace("/", "_").replace(" ", "_") if result.model_used else "unknown"
    source_name = Path(result.source_file).stem
    filename = f"{source_name}_{timestamp}_{safe_model}.json"

    path = out / filename

    with open(path, "w", encoding="utf-8") as f:
        f.write(result.model_dump_json(indent=2))

    return path


def _normalize_cut_text(value: str | None) -> str:
    return " ".join((value or "").split()).casefold()


def _panel_cut_key(panel: Panel) -> tuple[object, ...]:
    assert panel.length_mm is not None
    assert panel.width_mm is not None
    assert panel.thickness_mm is not None
    edge = panel.edge_banding
    return (
        panel.length_mm,
        panel.width_mm,
        panel.thickness_mm,
        _normalize_cut_text(panel.board_type),
        _normalize_cut_text(panel.material_code),
        _normalize_cut_text(edge.top),
        _normalize_cut_text(edge.right),
        _normalize_cut_text(edge.bottom),
        _normalize_cut_text(edge.left),
        _normalize_cut_text(panel.grain_direction),
    )


def _pricing_cut_key(panel: Panel) -> tuple[object, ...]:
    edge = panel.edge_banding
    return (
        _normalize_cut_text(panel.name),
        panel.length_mm,
        panel.width_mm,
        panel.thickness_mm,
        _normalize_cut_text(panel.board_type),
        _normalize_cut_text(panel.material_code),
        _normalize_cut_text(edge.top),
        _normalize_cut_text(edge.right),
        _normalize_cut_text(edge.bottom),
        _normalize_cut_text(edge.left),
        _normalize_cut_text(panel.grain_direction),
    )


def _is_pricing_candidate(panel: Panel) -> bool:
    return bool(
        (panel.name or "").strip()
        and panel.length_mm is not None
        and panel.length_mm > 0
        and panel.width_mm is not None
        and panel.width_mm > 0
        and ((panel.board_type or "").strip() or (panel.material_code or "").strip())
        and panel.qty > 0
    )


def _aggregate_confidence(grouped_panels: list[tuple[ExtractionResult, Panel]]) -> ConfidenceLevel:
    confidences = {panel.confidence for _, panel in grouped_panels}
    if ConfidenceLevel.LOW in confidences:
        return ConfidenceLevel.LOW
    if ConfidenceLevel.MEDIUM in confidences:
        return ConfidenceLevel.MEDIUM
    return ConfidenceLevel.HIGH


def _source_reference(result: ExtractionResult, panel: Panel) -> str:
    filename = Path(result.source_file).name
    if panel.part_id:
        return f"{filename}#{panel.part_id}"
    if panel.name:
        return f"{filename}:{panel.name}"
    return filename


def _pricing_reasons(grouped_panels: list[tuple[ExtractionResult, Panel]]) -> list[str]:
    if _aggregate_confidence(grouped_panels) == ConfidenceLevel.LOW:
        return ["LOW confidence"]
    return []


def write_xlsx(results: list[ExtractionResult], output_path: str) -> Path:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    from src.schema.extraction import SCALAR_FIELDS

    wb = Workbook()

    # ── Summary sheet (one row per drawing) ────────────────────
    ws = wb.active
    ws.title = "Extraction Results"

    headers = ["Source File", "Drawing Type", "Model", "Passes"]
    for fname in SCALAR_FIELDS:
        headers.extend([f"{fname.title()} (Value)", f"{fname.title()} (Confidence)"])
    headers.extend(["Panels", "Hardware", "Needs Review", "Tokens"])
    ws.append(headers)

    for result in results:
        row = [
            Path(result.source_file).name,
            result.drawing_type.value,
            result.model_used,
            result.passes_completed,
        ]
        for fname in SCALAR_FIELDS:
            f = getattr(result, fname, None)
            if f:
                row.extend([f.value or "", f.confidence.value])
            else:
                row.extend(["", ""])
        row.extend(
            [
                len(result.panels),
                len(result.hardware),
                "YES" if result.needs_review else "NO",
                result.tokens_used,
            ]
        )
        ws.append(row)

    pricing_groups: dict[tuple[object, ...], list[tuple[ExtractionResult, Panel]]] = {}
    production_groups: dict[tuple[object, ...], list[tuple[ExtractionResult, Panel]]] = {}
    review_panels: list[tuple[ExtractionResult, Panel, list[str]]] = []

    # ── Panels sheet (one row per panel — the CNC cutting list) ──
    ws_panels = wb.create_sheet("Cutting List")
    ws_panels.append(
        [
            "Source File",
            "Part ID",
            "Name",
            "Length (mm)",
            "Width (mm)",
            "Thickness (mm)",
            "Board Type",
            "Material Code",
            "Edge T",
            "Edge R",
            "Edge B",
            "Edge L",
            "Grain",
            "Qty",
            "Confidence",
            "Notes",
        ]
    )
    for result in results:
        for panel in result.panels:
            eb = panel.edge_banding
            ws_panels.append(
                [
                    Path(result.source_file).name,
                    panel.part_id or "",
                    panel.name or "",
                    panel.length_mm or "",
                    panel.width_mm or "",
                    panel.thickness_mm or "",
                    panel.board_type or "",
                    panel.material_code or "",
                    eb.top or "",
                    eb.right or "",
                    eb.bottom or "",
                    eb.left or "",
                    panel.grain_direction or "",
                    panel.qty,
                    panel.confidence.value,
                    panel.notes,
                ]
            )
            if _is_pricing_candidate(panel):
                pricing_groups.setdefault(_pricing_cut_key(panel), []).append((result, panel))
            reasons = panel_review_reasons(panel)
            if reasons:
                review_panels.append((result, panel, reasons))
            else:
                production_groups.setdefault(_panel_cut_key(panel), []).append((result, panel))

    # ── Hardware sheet ──────────────────────────────────────────
    ws_hw = wb.create_sheet("Hardware")
    ws_hw.append(["Source File", "Name", "Code", "Qty", "Notes"])
    for result in results:
        for hw in result.hardware:
            ws_hw.append(
                [
                    Path(result.source_file).name,
                    hw.name or "",
                    hw.code or "",
                    hw.qty,
                    hw.notes,
                ]
            )

    # ── Geometry-first inventory ───────────────────────────────
    ws_geometry = wb.create_sheet("Geometry Inventory")
    ws_geometry.append(
        [
            "Source File",
            "Group",
            "Category",
            "Length (mm)",
            "Width (mm)",
            "Thickness (mm)",
            "Source Panel Indices",
            "Observed Qty",
            "Minimum Qty",
            "Maximum Qty",
            "Confidence",
            "Notes",
        ]
    )
    for result in results:
        for group in result.geometry_groups:
            ws_geometry.append(
                [
                    Path(result.source_file).name,
                    group.group_id,
                    group.category,
                    group.length_mm if group.length_mm is not None else "",
                    group.width_mm if group.width_mm is not None else "",
                    group.thickness_mm if group.thickness_mm is not None else "",
                    ", ".join(str(index) for index in group.source_panel_indices),
                    group.observed_qty,
                    group.minimum_qty,
                    group.maximum_qty,
                    group.confidence.value,
                    group.notes,
                ]
            )

    # ── Raw panel quantities for pricing ──────────────────────
    ws_pricing = wb.create_sheet("Pricing Overview")
    ws_pricing.append(
        [
            "Part Name",
            "Material Code",
            "Board Type",
            "Thickness (mm)",
            "Length (mm)",
            "Width (mm)",
            "Edge T",
            "Edge R",
            "Edge B",
            "Edge L",
            "Grain",
            "Total Qty",
            "Net Area (m²)",
            "Confidence",
            "Pricing Status",
            "Projects",
            "Cabinets",
            "Source References",
        ]
    )
    for _, grouped_panels in sorted(
        pricing_groups.items(),
        key=lambda item: (
            _normalize_cut_text(item[1][0][1].name),
            _normalize_cut_text(item[1][0][1].material_code),
            _normalize_cut_text(item[1][0][1].board_type),
            item[1][0][1].thickness_mm is None,
            item[1][0][1].thickness_mm or 0,
            item[1][0][1].length_mm or 0,
            item[1][0][1].width_mm or 0,
            _normalize_cut_text(item[1][0][1].edge_banding.top),
            _normalize_cut_text(item[1][0][1].edge_banding.right),
            _normalize_cut_text(item[1][0][1].edge_banding.bottom),
            _normalize_cut_text(item[1][0][1].edge_banding.left),
            _normalize_cut_text(item[1][0][1].grain_direction),
        ),
    ):
        first_panel = grouped_panels[0][1]
        total_qty = sum(panel.qty for _, panel in grouped_panels)
        net_area = (
            sum(
                panel.length_mm * panel.width_mm * panel.qty
                for _, panel in grouped_panels
                if panel.length_mm is not None and panel.width_mm is not None
            )
            / 1_000_000
        )
        edge = first_panel.edge_banding
        reasons = _pricing_reasons(grouped_panels)
        projects = sorted({result.project_name.value for result, _ in grouped_panels if result.project_name.value})
        cabinets = sorted({result.cabinet_id.value for result, _ in grouped_panels if result.cabinet_id.value})
        source_references = sorted({_source_reference(result, panel) for result, panel in grouped_panels})
        ws_pricing.append(
            [
                first_panel.name or "",
                first_panel.material_code or "",
                first_panel.board_type or "",
                first_panel.thickness_mm if first_panel.thickness_mm is not None else "",
                first_panel.length_mm,
                first_panel.width_mm,
                edge.top or "",
                edge.right or "",
                edge.bottom or "",
                edge.left or "",
                first_panel.grain_direction or "",
                total_qty,
                round(net_area, 4),
                _aggregate_confidence(grouped_panels).value,
                "; ".join(reasons) if reasons else "READY TO PRICE",
                "; ".join(projects),
                "; ".join(cabinets),
                "; ".join(source_references),
            ]
        )

    # ── Consolidated production cutting list ───────────────────
    ws_production = wb.create_sheet("Production Cut List")
    ws_production.append(
        [
            "Material Code",
            "Board Type",
            "Thickness (mm)",
            "Length (mm)",
            "Width (mm)",
            "Edge T",
            "Edge R",
            "Edge B",
            "Edge L",
            "Grain",
            "Total Qty",
            "Net Area (m²)",
            "Confidence",
            "Part Names",
            "Projects",
            "Cabinets",
            "Source References",
        ]
    )
    production_totals: list[tuple[Panel, int, float]] = []
    for _, grouped_panels in sorted(
        production_groups.items(),
        key=lambda item: (
            _normalize_cut_text(item[1][0][1].material_code),
            _normalize_cut_text(item[1][0][1].board_type),
            item[1][0][1].thickness_mm or 0,
            item[1][0][1].length_mm or 0,
            item[1][0][1].width_mm or 0,
            _normalize_cut_text(item[1][0][1].edge_banding.top),
            _normalize_cut_text(item[1][0][1].edge_banding.right),
            _normalize_cut_text(item[1][0][1].edge_banding.bottom),
            _normalize_cut_text(item[1][0][1].edge_banding.left),
            _normalize_cut_text(item[1][0][1].grain_direction),
        ),
    ):
        first_panel = grouped_panels[0][1]
        total_qty = sum(panel.qty for _, panel in grouped_panels)
        net_area = 0.0
        for _, panel in grouped_panels:
            assert panel.length_mm is not None
            assert panel.width_mm is not None
            net_area += panel.length_mm * panel.width_mm * panel.qty
        net_area /= 1_000_000
        production_totals.append((first_panel, total_qty, net_area))

        part_names = sorted({panel.name for _, panel in grouped_panels if panel.name})
        projects = sorted({result.project_name.value for result, _ in grouped_panels if result.project_name.value})
        cabinets = sorted({result.cabinet_id.value for result, _ in grouped_panels if result.cabinet_id.value})
        source_references = sorted({_source_reference(result, panel) for result, panel in grouped_panels})
        edge = first_panel.edge_banding
        confidence = _aggregate_confidence(grouped_panels)
        ws_production.append(
            [
                first_panel.material_code or "",
                first_panel.board_type or "",
                first_panel.thickness_mm,
                first_panel.length_mm,
                first_panel.width_mm,
                edge.top or "",
                edge.right or "",
                edge.bottom or "",
                edge.left or "",
                first_panel.grain_direction or "",
                total_qty,
                round(net_area, 4),
                confidence.value,
                "; ".join(part_names),
                "; ".join(projects),
                "; ".join(cabinets),
                "; ".join(source_references),
            ]
        )

    # ── Net material demand by material and thickness ──────────
    ws_materials = wb.create_sheet("Material Summary")
    ws_materials.append(
        [
            "Material Code",
            "Board Type",
            "Thickness (mm)",
            "Total Qty",
            "Distinct Cut Specs",
            "Net Area (m²)",
        ]
    )
    material_groups: dict[tuple[str, str, float], list[tuple[Panel, int, float]]] = {}
    for panel, total_qty, net_area in production_totals:
        assert panel.thickness_mm is not None
        material_key = (
            _normalize_cut_text(panel.material_code),
            _normalize_cut_text(panel.board_type),
            float(panel.thickness_mm),
        )
        material_groups.setdefault(material_key, []).append((panel, total_qty, net_area))

    for _, material_rows in sorted(material_groups.items()):
        first_panel = material_rows[0][0]
        ws_materials.append(
            [
                first_panel.material_code or "",
                first_panel.board_type or "",
                first_panel.thickness_mm,
                sum(total_qty for _, total_qty, _ in material_rows),
                len(material_rows),
                round(sum(net_area for _, _, net_area in material_rows), 4),
            ]
        )

    # ── Panels excluded from production totals ─────────────────
    ws_review = wb.create_sheet("Panel Review")
    ws_review.append(
        [
            "Source File",
            "Project",
            "Cabinet",
            "Part ID",
            "Name",
            "Length (mm)",
            "Width (mm)",
            "Thickness (mm)",
            "Board Type",
            "Material Code",
            "Edge T",
            "Edge R",
            "Edge B",
            "Edge L",
            "Grain",
            "Qty",
            "Confidence",
            "Review Reason",
            "Notes",
        ]
    )
    for result, panel, reasons in review_panels:
        edge = panel.edge_banding
        ws_review.append(
            [
                Path(result.source_file).name,
                result.project_name.value or "",
                result.cabinet_id.value or "",
                panel.part_id or "",
                panel.name or "",
                panel.length_mm if panel.length_mm is not None else "",
                panel.width_mm if panel.width_mm is not None else "",
                panel.thickness_mm if panel.thickness_mm is not None else "",
                panel.board_type or "",
                panel.material_code or "",
                edge.top or "",
                edge.right or "",
                edge.bottom or "",
                edge.left or "",
                panel.grain_direction or "",
                panel.qty,
                panel.confidence.value,
                "; ".join(reasons),
                panel.notes,
            ]
        )

    wb.save(str(path))
    return path


def write_token_report(
    metrics: list[DrawingMetrics],
    output_path: str,
    schema_overhead: SchemaOverhead | None = None,
) -> Path:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    # ── Summary sheet ────────────────────────────────────────
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary.append(
        [
            "Filename",
            "Drawing Type",
            "Total Tokens",
            "Peak Input Tokens",
            "Overflow",
            "Context Window",
            "Passes Completed",
        ]
    )
    for m in metrics:
        peak = max((p.peak_input_tokens for p in m.passes), default=0)
        ws_summary.append(
            [
                m.filename,
                m.drawing_type,
                m.total_tokens,
                peak,
                "YES" if m.overflow else "NO",
                m.context_window,
                len(m.passes),
            ]
        )

    # ── Per-Pass sheet ───────────────────────────────────────
    ws_passes = wb.create_sheet("Per-Pass")
    ws_passes.append(
        [
            "Filename",
            "Pass",
            "Requests",
            "Input Tokens",
            "Output Tokens",
            "Peak Input Tokens",
            "Est. Image Tokens",
            "Overflow",
            "Overflow Tokens",
        ]
    )
    for m in metrics:
        for p in m.passes:
            ws_passes.append(
                [
                    m.filename,
                    p.pass_name,
                    p.request_count,
                    p.total_input_tokens,
                    p.total_output_tokens,
                    p.peak_input_tokens,
                    p.estimated_image_tokens,
                    "YES" if p.overflow else "NO",
                    p.overflow_tokens,
                ]
            )

    # ── Schema Overhead sheet ────────────────────────────────
    if schema_overhead is not None:
        ws_schema = wb.create_sheet("Schema Overhead")
        ws_schema.append(["Component", "Estimated Tokens"])
        ws_schema.append(["Output JSON Schema", schema_overhead.output_schema_tokens])
        ws_schema.append(["Tool Definitions", schema_overhead.tool_schemas_tokens])
        ws_schema.append(["System Prompt (longest)", schema_overhead.system_prompt_tokens])
        ws_schema.append(["Total Fixed Overhead", schema_overhead.total_fixed_tokens])
        ws_schema.append([])
        ws_schema.append(["Per-Pass Prompt Tokens", ""])
        for name, tokens in schema_overhead.per_pass_prompts.items():
            ws_schema.append([name, tokens])

    wb.save(str(path))
    return path
